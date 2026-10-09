import ast
import asyncio
import math
import re
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt


def _load():
    src = Path(__file__).resolve().parents[1].joinpath("main.py").read_text()
    names = ["_SENTENCE_ABBREVIATIONS", "_review_two_space_style", "_review_normalise_sentence_spacing", "_review_text_to_html", "_review_enforce_indent", "_human_review_spelling_notes", "_human_review_candidate_terms", "_human_review_missing_research_terms", "_human_research_parse_findings", "_human_research_unverified_terms", "_human_ai_agent_research_footer", "_human_worker_feedback", "_review_strip_model_footer", "_review_split_output"]
    chunks = []
    for name in names:
        start = src.index(name + " =") if name.startswith("_SENT") else src.index("def " + name)
        nxt = src.index("\n\n\n", start)
        chunks.append(src[start:nxt])
    ns = {"re": re, "math": math, "_parse_json_object": __import__("json").loads}
    exec("\n\n".join(chunks), ns)
    return ns


NS = _load()


class ReviewHelpers(unittest.TestCase):
    def test_spacing(self):
        out = NS["_review_normalise_sentence_spacing"]("He left. She stayed. Ms. Lee spoke. J. Smith agreed.")
        self.assertEqual(out, "He left.  She stayed.  Ms. Lee spoke.  J. Smith agreed.")

    def test_missing_research_terms_need_structured_confident_findings(self):
        missing = NS["_human_review_missing_research_terms"]
        candidates = ["Kneeland", "Tyshawn", "Family to Family Program"]
        raw_queries = "ACTUAL GOOGLE SEARCH QUERIES:\n- Kneeland Tyshawn Family to Family Program\nACTUAL SEARCH SOURCES:\n- example.org"
        self.assertEqual(missing(candidates, raw_queries), candidates)
        grounded = (
            "Kneeland | Kneeland | Name spelling confirmed. | confidence: yes\n"
            "Tyshawn | Tyshawn | Name spelling unclear. | confidence: no"
        )
        self.assertEqual(missing(candidates, grounded), ["Tyshawn", "Family to Family Program"])

    def test_two_space_detect(self):
        self.assertTrue(NS["_review_two_space_style"](["One.  Two.  Three."]))
        self.assertFalse(NS["_review_two_space_style"](["One. Two. Three."]))

    def test_indent_enforced(self):
        src = ["\tFirst para.\n\n\tSecond para."]
        out = NS["_review_enforce_indent"]("First para.\n\nHEADING\n\nSecond para.", src)
        self.assertEqual(out, "\tFirst para.\n\n\nHEADING\n\n\tSecond para.".replace("\n\n\nHEADING", "\n\nHEADING"))

    def test_indent_applied_even_when_parts_lost_tabs(self):
        out = NS["_review_enforce_indent"]("First para.\n\nSecond para.\n\nClient spellings: Ann.", ["First para.\n\nSecond para."])
        self.assertEqual(out, "\tFirst para.\n\n\tSecond para.\n\nClient spellings: Ann.")

    def test_split(self):
        text, data = NS["_review_split_output"]('<<<TRANSCRIPT>>>\n\tHello.  World.\n<<<NOTES>>>\n{"summary": "ok"}')
        self.assertEqual(text, "\tHello.  World.")
        self.assertEqual(data["summary"], "ok")

    def test_batch_footer_is_removed_without_losing_transcript(self):
        strip_footer = NS["_review_strip_model_footer"]
        first = strip_footer("Part one text.\nClient spellings: Ann.\nResearch Notes:\n- Note for first part.")
        second = strip_footer("Part two text.\nClient spellings: Lee.\nResearch Notes:\n- Note for second part.")
        self.assertEqual(first, "Part one text.")
        self.assertEqual(second, "Part two text.")
        self.assertEqual("\n\n".join([first, second]), "Part one text.\n\nPart two text.")

    def test_cross_part_client_spellings_and_worker_research_are_collected(self):
        notes = NS["_human_review_spelling_notes"]([
            {"label": "Part 1", "text": "Transcript body.\nClient spellings: Jon, My spellings: Renee"},
            {"label": "Part 3", "text": "Transcript body.\nClient spellings: John; I searched: Summit Psych\nResearch Notes:\nSummit Psych is the provider named in the recording."},
        ])
        self.assertIn("CONSOLIDATED CLIENT SPELLINGS FROM ALL PARTS", notes)
        self.assertIn("Jon, John", notes)
        self.assertIn("Part 1: Jon", notes)
        self.assertIn("Part 3: John", notes)
        self.assertIn("WORKER-REPORTED SEARCH TERMS FROM ALL PARTS", notes)
        self.assertIn("Part 3: Summit Psych", notes)
        self.assertIn("provider named in the recording", notes)
        self.assertNotIn("Renee", notes)

    def test_no_worker_spelling_notes_returns_empty_context(self):
        self.assertEqual(NS["_human_review_spelling_notes"]([{"text": "Ordinary transcript text."}]), "")

    def test_candidate_terms_exclude_client_spellings_roles_and_contact_handles(self):
        candidates = NS["_human_review_candidate_terms"](
            [{"text": "Ms. Kneeland spoke with Kyra Kneeland. kyrakneeland55@gmail.com Kyra.Kneeland55.\nClient spellings: Kneeland, Kyra, Zaden; My spellings: None; I researched: Kneeland, Caseworker."}],
            "Kneeland spoke with Caseworker.",
            "",
        )
        self.assertFalse({"Kneeland", "Kyra", "Kyra Kneeland", "Zaden", "Caseworker", "kyrakneeland55", "Kyra.Kneeland55"} & set(candidates))
        self.assertIn("Family-to-Family Program", NS["_human_review_candidate_terms"](
            [{"text": "The Family-to-Family Program was mentioned."}], "", ""
        ))

    def test_worker_sees_only_own_feedback_as_admin(self):
        feedback = NS["_human_worker_feedback"]({
            "worker_uid": "w1",
            "admin_feedback": "Overall comment",
            "worker_rating": 3,
            "segments": [{"id": "part-a", "worker_uid": "w1"}, {"id": "part-b", "worker_uid": "w2"}],
            "part_ratings": {
                "part-a": {"worker_uid": "w1", "rating": 3, "note": "AI review: Verify the organization name.", "source": "ai"},
                "part-b": {"worker_uid": "w2", "rating": 2, "note": "Other worker note", "source": "admin"},
            },
        }, "w1")
        self.assertEqual(feedback, [{
            "label": "Part 1", "rating": 3.0,
            "note": "Verify the organization name.", "rater": "Admin",
        }])
        self.assertNotIn("source", feedback[0])
        self.assertNotIn("AI", feedback[0]["note"])

    def test_legacy_whole_job_feedback_is_included_for_assigned_worker(self):
        feedback = NS["_human_worker_feedback"]({
            "worker_uid": "w1", "worker_rating": 4,
            "admin_feedback": "Clear work; please verify names.",
        }, "w1")
        self.assertEqual(feedback[0]["label"], "Overall job review")
        self.assertEqual(feedback[0]["note"], "Clear work; please verify names.")
        self.assertEqual(feedback[0]["rater"], "Admin")

    def test_html(self):
        self.assertEqual(NS["_review_text_to_html"]("a<b\n\nc"), "<div>a&lt;b</div><div><br></div><div>c</div>")


class ProofreaderLabelSerialization(unittest.TestCase):
    def test_proofreader_parts_use_position_labels_without_ai_attribution(self):
        source = Path(__file__).resolve().parents[1].joinpath("main.py").read_text()
        serializer = source[source.index("def _human_public_for("):source.index("def _human_available_public_for(")]
        self.assertIn('"author_label": f"Worker {index + 1}"', serializer)
        self.assertNotIn("AI draft:", serializer)


class WholeJobAiTakeover(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source_path = Path(__file__).resolve().parents[1].joinpath("main.py")
        cls.source = source_path.read_text()
        tree = ast.parse(cls.source)
        names = {"_human_is_split_job", "_human_whole_job_ai_takeover_allowed"}
        functions = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in names]
        cls.helpers = {}
        exec(compile(ast.Module(body=functions, type_ignores=[]), str(source_path), "exec"), cls.helpers)
        cls.routes = {node.name: node for node in tree.body if isinstance(node, ast.AsyncFunctionDef)}

    def test_takeover_is_allowed_only_before_any_part_or_proofreader_starts(self):
        allowed = self.helpers["_human_whole_job_ai_takeover_allowed"]
        base = {"split_mode": "dual", "status": "split_assigned", "segments": [
            {"id": "a", "status": "available"}, {"id": "b", "status": "approved"},
        ]}
        self.assertTrue(allowed(base))
        claimed = {**base, "segments": [dict(base["segments"][0]), {"id": "b", "status": "in_progress", "worker_uid": "worker-2"}]}
        self.assertFalse(allowed(claimed))
        submitted = {**base, "segments": [dict(base["segments"][0]), {"id": "b", "status": "submitted"}]}
        self.assertFalse(allowed(submitted))
        proofreader = {**base, "proofreader_status": "assigned"}
        self.assertFalse(allowed(proofreader))
        self.assertFalse(allowed({"split_mode": "single", "segments": base["segments"]}))

    def test_route_stores_paused_parts_and_restores_them_on_failure(self):
        route_source = ast.unparse(self.routes["human_admin_assign_ai_agent"])
        runner_source = ast.unparse(self.routes["_human_run_ai_agent"])
        self.assertIn("_require_ai_agent_assignment", route_source)
        self.assertIn("_human_whole_job_ai_takeover_allowed", route_source)
        self.assertIn("ai_agent_paused_segments", route_source)
        self.assertIn("ai_agent_previous_split_mode", route_source)
        self.assertIn("ai_agent_paused_segments", runner_source)
        self.assertIn("restored_segments", runner_source)
        self.assertIn("firestore.DELETE_FIELD", runner_source)

    def test_worker_and_client_serialization_redacts_agent_identity(self):
        serializer = self.source[self.source.index("def _human_public_for("):self.source.index("def _human_available_public_for(")]
        self.assertIn('"ai_agent_name"', serializer)
        self.assertIn('"ai_agent_model_ids"', serializer)
        self.assertIn('"ai_agent_docx"', serializer)
        self.assertIn('"author_label": f"Worker {index + 1}"', serializer)


class TemplateDocxRendering(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source_path = Path(__file__).resolve().parents[1].joinpath("main.py")
        tree = ast.parse(source_path.read_text())
        names = {"_human_template_attachment", "_human_template_docx_profile", "_human_template_line_kind", "_human_template_source_kind", "_human_template_render_docx"}
        functions = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in names]
        cls.helpers = {"Document": Document, "BytesIO": BytesIO, "deepcopy": deepcopy, "re": re}
        exec(compile(ast.Module(body=functions, type_ignores=[]), str(source_path), "exec"), cls.helpers)

    def test_template_selection_never_guesses_between_word_files(self):
        choose = self.helpers["_human_template_attachment"]
        selected = choose({"instruction_attachments": [{"name": "notes.pdf"}, {"name": "letter.docx", "storage_path": "one"}]})
        self.assertEqual(selected[1]["storage_path"], "one")
        with self.assertRaisesRegex(ValueError, "More than one Word document"):
            choose({"instruction_attachments": [{"name": "one.docx"}, {"name": "two.docx"}]})
        selected = choose({"instruction_attachments": [
            {"name": "one.docx"}, {"name": "two.docx", "purpose": "template"},
        ]})
        self.assertEqual(selected[1]["name"], "two.docx")

    def test_render_keeps_the_attached_template_page_and_paragraph_formatting(self):
        template = Document()
        section = template.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
        template.styles["Normal"].font.name = "Times New Roman"
        template.styles["Normal"].font.size = Pt(12)
        heading = template.add_paragraph()
        heading.alignment = 1
        heading.add_run("EXAMPLE TITLE").bold = True
        template.add_paragraph()
        body = template.add_paragraph()
        body.paragraph_format.line_spacing = 1
        body.paragraph_format.space_before = Pt(0)
        body.paragraph_format.space_after = Pt(0)
        body.add_run("\tExample body placeholder.")
        template.add_paragraph("Client spellings: Example.")
        original = BytesIO()
        template.save(original)

        profile = self.helpers["_human_template_docx_profile"](original.getvalue())
        result_bytes = self.helpers["_human_template_render_docx"](
            original.getvalue(), "EXAMPLE TITLE\n\n\tBody paragraph.\n\nClient spellings: Jon."
        )
        result = Document(BytesIO(result_bytes))
        self.assertIn("Page size: 8.50 x 11.00 inches", profile)
        self.assertEqual([paragraph.text for paragraph in result.paragraphs], ["EXAMPLE TITLE", "", "\tBody paragraph.", "", "Client spellings: Jon."])
        self.assertEqual(result.sections[0].top_margin, section.top_margin)
        self.assertEqual(result.sections[0].page_width, section.page_width)
        self.assertEqual(result.styles["Normal"].font.name, "Times New Roman")
        self.assertEqual(result.styles["Normal"].font.size, Pt(12))
        self.assertEqual(result.paragraphs[0].alignment, 1)
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertEqual(result.paragraphs[2].paragraph_format.line_spacing, 1)
        self.assertEqual(result.paragraphs[2].paragraph_format.space_after, Pt(0))

    def test_template_indent_is_not_doubled_with_default_transcript_tabs(self):
        template = Document()
        body = template.add_paragraph("Original body paragraph.")
        body.paragraph_format.first_line_indent = Inches(0.5)
        original = BytesIO()
        template.save(original)
        result_bytes = self.helpers["_human_template_render_docx"](original.getvalue(), "\tReplacement body paragraph.")
        result = Document(BytesIO(result_bytes))
        self.assertEqual(result.paragraphs[0].text, "Replacement body paragraph.")
        self.assertEqual(result.paragraphs[0].paragraph_format.first_line_indent, Inches(0.5))

    def test_table_templates_fail_closed_without_creating_a_misleading_docx(self):
        template = Document()
        template.add_table(rows=1, cols=1).cell(0, 0).text = "Name:"
        output = BytesIO()
        template.save(output)
        with self.assertRaisesRegex(ValueError, "uses tables"):
            self.helpers["_human_template_docx_profile"](output.getvalue())
        with self.assertRaisesRegex(ValueError, "uses tables"):
            self.helpers["_human_template_render_docx"](output.getvalue(), "Draft text")

    def test_assignment_uses_the_job_template_and_only_full_admin_can_download(self):
        source_path = Path(__file__).resolve().parents[1].joinpath("main.py")
        routes = {node.name: node for node in ast.parse(source_path.read_text()).body if isinstance(node, ast.AsyncFunctionDef)}
        assignment = ast.unparse(routes["human_admin_assign_ai_agent"])
        runner = ast.unparse(routes["_human_run_ai_agent"])
        download = ast.unparse(routes["human_admin_ai_agent_template_docx"])
        self.assertIn("_human_template_docx_bytes", assignment)
        self.assertIn("_human_template_docx_profile", assignment)
        self.assertIn("_human_template_render_docx", runner)
        self.assertIn("_human_store_raw_bytes", runner)
        self.assertIn("_require_ai_agent_assignment", download)
        self.assertIn("human-workflow/{job_id}/ai-drafts/", download)


class AiAgentAuthorization(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source = Path(__file__).resolve().parents[1].joinpath("main.py").read_text()
        cls.functions = {
            node.name: node for node in ast.parse(source).body
            if isinstance(node, ast.AsyncFunctionDef)
        }

    def test_catalog_and_assignment_require_full_admin(self):
        for name in ("human_admin_ai_agents", "human_admin_assign_ai_agent"):
            with self.subTest(route=name):
                calls = {
                    node.func.id for node in ast.walk(self.functions[name])
                    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
                }
                self.assertIn("_require_ai_agent_assignment", calls)
                self.assertNotIn("_require_human_job_admin", calls)


if __name__ == "__main__":
    unittest.main()

class AiAgentCatalog(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source = Path(__file__).resolve().parents[1].joinpath("main.py").read_text()
        tree = __import__("ast").parse(source)
        node = next(n for n in tree.body if isinstance(n, __import__("ast").Assign) and any(getattr(t, "id", None) == "HUMAN_AI_AGENTS" for t in n.targets))
        cls.agents = __import__("ast").literal_eval(node.value)

    def test_agent_catalog_and_requested_model_pairs(self):
        self.assertEqual(set(self.agents), {
            "general-gpt", "template-claude", "pdf-gemini", "text-messages-gemini", "letter-opus",
        })
        expected_audio_models = ["gpt-5.6-sol", "claude-opus-5-5"]
        self.assertEqual(self.agents["general-gpt"]["models"], ["gemini-3.8-flash", "gpt-5.6-luna"])
        self.assertEqual(self.agents["template-claude"]["models"], expected_audio_models)
        image_models = ["gemini-3.5-flash-lite", "gemini-3.8-flash"]
        self.assertEqual(self.agents["pdf-gemini"]["models"], image_models)
        self.assertEqual(self.agents["text-messages-gemini"]["models"], image_models)
        self.assertEqual(self.agents["letter-opus"]["models"], expected_audio_models)

    def test_agents_are_internal_not_email_accounts(self):
        for agent in self.agents.values():
            self.assertNotIn("email", agent)
            self.assertNotIn("password", agent)


class HealthEndpointPrivacy(unittest.TestCase):
    def test_public_health_endpoint_returns_status_only(self):
        source_path = Path(__file__).resolve().parents[1].joinpath("main.py")
        tree = ast.parse(source_path.read_text())
        function = next(node for node in tree.body if isinstance(node, ast.AsyncFunctionDef) and node.name == "health_check")
        function.decorator_list = []
        isolated = ast.Module(body=[function], type_ignores=[])
        namespace = {}
        exec(compile(isolated, str(source_path), "exec"), namespace)
        self.assertEqual(asyncio.run(namespace["health_check"]()), {"status": "healthy"})


class AiModelRouting(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source_path = Path(__file__).resolve().parents[1].joinpath("main.py")
        cls.source = source_path.read_text(encoding="utf-8")
        cls.tree = ast.parse(cls.source)
        cls.functions = {
            node.name: node for node in cls.tree.body
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
        }
        cls.assignments = {
            target.id: ast.literal_eval(node.value)
            for node in cls.tree.body if isinstance(node, ast.Assign)
            for target in node.targets if isinstance(target, ast.Name)
            and target.id in {"AI_REVIEW_MODEL_CHAIN", "HUMAN_ADMIN_PROOFREAD_MODEL_CHAIN", "HUMAN_AUDIO_AGENT_MODEL_CHAIN", "HUMAN_GENERAL_AGENT_MODEL_CHAIN", "HUMAN_PDF_AGENT_MODEL_CHAIN", "HUMAN_TEXT_MESSAGES_MODEL_CHAIN", "WORKER_DRAFT_FORMAT_MODEL_CHAIN", "WORKER_DRAFT_PROOFREAD_MODEL_CHAIN", "ASK_MODEL_CATALOGUE"}
        }

    def test_ai_review_preserves_emphatic_repetition_and_limits_self_corrections(self):
        node = next(
            item for item in self.tree.body
            if isinstance(item, ast.Assign)
            and any(isinstance(target, ast.Name) and target.id == "_REVIEW_SYSTEM" for target in item.targets)
        )
        prompt = ast.literal_eval(node.value)
        self.assertIn("immediately and unmistakably corrects that same wording", prompt)
        self.assertIn("very, very hot", prompt)
        self.assertIn("Never treat emphasis", prompt)
        self.assertIn("application will append one consolidated client-spellings line", prompt)
        self.assertIn("Do not add a closing spellings line or `Research Notes:` section to an individual batch", prompt)

    def test_requested_model_chains_are_primary_then_fallback(self):
        self.assertEqual(self.assignments["AI_REVIEW_MODEL_CHAIN"], (("gpt-5.6-terra", "openai"), ("gemini-3.8-flash", "gemini"), ("gpt-5.6-luna", "openai")))
        self.assertEqual(self.assignments["HUMAN_ADMIN_PROOFREAD_MODEL_CHAIN"], (("claude-sonnet-5-5", "claude"), ("gemini-3.8-flash", "gemini"), ("gpt-5.6-luna", "openai")))
        self.assertEqual(self.assignments["HUMAN_AUDIO_AGENT_MODEL_CHAIN"], (("gpt-5.6-sol", "openai"), ("claude-opus-5-5", "claude")))
        self.assertEqual(self.assignments["HUMAN_GENERAL_AGENT_MODEL_CHAIN"], (("gemini-3.8-flash", "gemini"), ("gpt-5.6-luna", "openai")))
        self.assertEqual(self.assignments["HUMAN_PDF_AGENT_MODEL_CHAIN"], (("gemini-3.5-flash-lite", "gemini"), ("gemini-3.8-flash", "gemini")))
        self.assertEqual(self.assignments["HUMAN_TEXT_MESSAGES_MODEL_CHAIN"], (("gemini-3.5-flash-lite", "gemini"), ("gemini-3.8-flash", "gemini")))
        self.assertEqual(self.assignments["WORKER_DRAFT_FORMAT_MODEL_CHAIN"], (("gpt-5.6-luna", "openai"), ("gemini-3.5-flash-lite", "gemini")))
        self.assertEqual(self.assignments["WORKER_DRAFT_PROOFREAD_MODEL_CHAIN"], (("claude-haiku-5-5", "claude"), ("gpt-5.6-luna", "openai")))
        flash_lite = next(model for model in self.assignments["ASK_MODEL_CATALOGUE"] if model["id"] == "gemini-3.5-flash-lite")
        self.assertEqual((flash_lite["provider"], flash_lite["tier"], flash_lite["credits"], flash_lite["transcript_only"]), ("gemini", "standard", 2, True))

    def test_audio_agents_compare_both_transcripts_and_fail_if_either_is_missing(self):
        transcriber = self.functions["_human_ai_transcribe_audio"]
        calls = {node.func.id for node in ast.walk(transcriber) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)}
        self.assertTrue({"transcribe_with_assemblyai", "transcribe_with_deepgram", "_human_ai_pair_asr_transcripts"}.issubset(calls))
        self.assertTrue(any(
            isinstance(node.func, ast.Attribute) and node.func.attr == "gather"
            and isinstance(node.func.value, ast.Name) and node.func.value.id == "asyncio"
            for node in ast.walk(transcriber) if isinstance(node, ast.Call)
        ))
        agent_source = "".join(node.value for node in ast.walk(self.functions["_human_ai_agent_generate"]) if isinstance(node, ast.Constant) and isinstance(node.value, str))
        self.assertIn("Compare the independent AssemblyAI and Deepgram transcripts", agent_source)
        self.assertIn("Compare both source transcripts", agent_source)

        helper = self.functions["_human_ai_pair_asr_transcripts"]
        namespace = {"asyncio": asyncio}
        exec(compile(ast.Module(body=[helper], type_ignores=[]), "main.py", "exec"), namespace)
        pair = namespace["_human_ai_pair_asr_transcripts"]
        result = pair({"status": "completed", "transcription": "Assembly words."}, {"status": "completed", "transcript": "Deepgram words."})
        self.assertEqual(result, {"AssemblyAI": "Assembly words.", "Deepgram": "Deepgram words."})
        with self.assertRaisesRegex(RuntimeError, "Both AssemblyAI and Deepgram"):
            pair({"status": "completed", "transcription": "Assembly words."}, {"status": "failed"})

    def test_ai_review_and_both_agent_passes_use_fallback_chain(self):
        review_calls = [node for node in ast.walk(self.functions["human_admin_ai_review"]) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "_human_call_model_chain"]
        self.assertEqual(len(review_calls), 1)
        self.assertTrue(any(isinstance(arg, ast.Name) and arg.id == "HUMAN_ADMIN_PROOFREAD_MODEL_CHAIN" for arg in review_calls[0].args))
        self.assertTrue(any(keyword.arg == "response_validator" and isinstance(keyword.value, ast.Name) and keyword.value.id == "_review_validate_output" for keyword in review_calls[0].keywords))
        self.assertNotIn("resolve_ask_model", {node.func.id for node in ast.walk(self.functions["human_admin_ai_review"]) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)})
        agent_calls = [node for node in ast.walk(self.functions["_human_ai_agent_generate"]) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "_human_call_model_chain"]
        audio_calls = [call for call in agent_calls if any(isinstance(arg, ast.Name) and arg.id == "audio_agent_chain" for arg in call.args)]
        image_calls = [call for call in agent_calls if any(isinstance(arg, ast.Name) and arg.id == "HUMAN_TEXT_MESSAGES_MODEL_CHAIN" for arg in call.args)]
        pdf_calls = [call for call in agent_calls if any(isinstance(arg, ast.Name) and arg.id == "HUMAN_PDF_AGENT_MODEL_CHAIN" for arg in call.args)]
        agent_source = ast.unparse(self.functions["_human_ai_agent_generate"])
        self.assertEqual(len(audio_calls), 2)
        self.assertEqual(len(image_calls), 1)
        self.assertEqual(len(pdf_calls), 2)
        self.assertEqual(len(agent_calls), 5)
        self.assertIn("HUMAN_GENERAL_AGENT_MODEL_CHAIN if agent_id == 'general-gpt'", agent_source)
        self.assertIn("HUMAN_PDF_AGENT_MODEL_CHAIN if job.get('pdf_review')", ast.unparse(self.functions["_human_image_review_draft"]))
        image_batch = ast.unparse(self.functions["_human_image_batch_compute"])
        self.assertIn("HUMAN_PDF_AGENT_MODEL_CHAIN", image_batch)
        letter_review = ast.unparse(self.functions["_human_run_letter_ai_review"])
        self.assertIn("AI_REVIEW_MODEL_CHAIN", letter_review)
        admin_review = ast.unparse(self.functions["human_admin_ai_review"])
        self.assertIn("source_parts=parts", admin_review)
        self.assertIn("_review_strip_model_footer", ast.unparse(self.functions["_review_validate_output"]))

    def test_research_footer_uses_only_grounded_terms_sources_and_explicit_no_search_note(self):
        footer = NS["_human_ai_agent_research_footer"]
        no_search = footer("Transcript body.", "")
        self.assertIn("Research Notes:\nNo external searches were needed for this transcript.", no_search)
        self.assertNotIn("I researched:", no_search)
        ungrounded = footer("Transcript body.\nI researched: Example Org.", "")
        self.assertNotIn("I researched:", ungrounded)
        self.assertIn("No external searches were needed for this transcript.", ungrounded)
        malformed = footer("Transcript body.\nI researched: Example Org.", "A result without structured findings.")
        self.assertTrue(malformed.startswith("Transcript body."))
        self.assertIn("Research Notes:\nOnline research did not confirm a reliable finding", malformed)
        unverified = footer(
            "Transcript body.",
            "UNVERIFIED TERMS (do not correct from memory): Walden Park, Summit Clinic",
        )
        self.assertIn("could not confidently verify these terms; they were left as dictated: Walden Park, Summit Clinic.", unverified)
        table = footer(
            "Transcript body about Family to Family.",
            "| Dictated form | Verified spelling | Meaning | Confident |\n| --- | --- | --- | --- |\n"
            "| Family to Family | Family to Family Program | A family-support initiative. | Yes |",
        )
        self.assertIn("I researched: Family to Family.", table)
        self.assertNotIn("Family to Family Program", table)
        self.assertIn("- Family to Family: A family-support initiative. (confidence: yes)", table)
        absent = footer("Transcript body.", "Zebra Corp | Zebra Corporation | A company never dictated here. | yes")
        self.assertNotIn("Zebra", absent)
        grounded = footer(
            "Transcript body about Example Org.\nClient spellings: Ann, My spellings: Lee.",
            "Example Org | Example Organization | A service named in the transcript. | yes\n\n"
            "ACTUAL GOOGLE SEARCH QUERIES:\n- Example Org official service\n\n"
            "ACTUAL SEARCH SOURCES:\n- Official site: https://example.org",
        )
        self.assertIn("My spellings: Lee; I researched: Example Org.", grounded)
        self.assertIn("Research Notes:\n- Example Org: A service named in the transcript. (confidence: yes)", grounded)
        self.assertNotIn("Actual Google searches", grounded)
        self.assertNotIn("Sources:", grounded)
        self.assertNotIn("https://example.org", grounded)

    def test_admin_final_footer_consolidates_part_spellings_and_unresolved_research(self):
        footer = NS["_human_ai_agent_research_footer"]
        transcript = "Woodward Park was the meeting place.\nClient spellings: ModelOnly."
        parts = [
            {"text": "Part one.\nClient spellings: Alia; My spellings: Renee."},
            {"text": "Part two.\nClient spellings: Zaliah; My spellings: Alia."},
        ]
        research = (
            "Woodward Park | Woodward Park | A local park named as the meeting place in the transcript. | confidence: yes\n"
            "UNVERIFIED TERMS (do not correct from memory): Summit Psych"
        )
        result = footer(transcript, research, source_parts=parts)
        self.assertIn("Client spellings: Alia, Zaliah; My spellings: Renee; I researched: Woodward Park.", result)
        self.assertNotIn("ModelOnly", result)
        self.assertEqual(result.count("Client spellings:"), 1)
        self.assertIn("Research Notes:\n- Woodward Park: A local park named as the meeting place in the transcript. (confidence: yes)", result)
        self.assertIn("Online research could not confidently verify: Summit Psych. These terms were left as dictated.", result)
        self.assertNotIn("UNVERIFIED TERMS", result)

    def test_candidate_terms_skip_sentence_start_words_and_titles(self):
        text = "Present at the home was Ms. Wade. Both children were home. They visited Woodward Park. Then Ohio came up. Historically Owens lived there. If Aaliyah came later."
        terms = NS["_human_review_candidate_terms"]([{"text": text}], "", "")
        self.assertIn("Woodward Park", terms)
        for junk in ("Present", "Both", "Historically", "If Aaliyah"):
            self.assertNotIn(junk, terms)
        self.assertFalse([term for term in terms if "Wade" in term and "Ms" in term])

    def test_research_footer_removes_duplicate_spellings_roles_and_contact_details(self):
        transcript = (
            "Transcript body about the Family to Family Program.\n"
            "Client spellings: Kneeland, Kyra, Zaden, Zaliah, Zaire, Tyshawn; "
            "My spellings: None; I researched: Kneeland, Kyra.Kneeland55, Zaden, Zaliah, Caseworker, Tyshawn.\n\n"
            "Research Notes:\n- Old model-written duplicate notes."
        )
        research = "\n".join([
            "Kneeland | Kneeland | Client's surname, already supplied by the client. | yes",
            "Kneeland | Kneeland | Repeated spelling note for the same surname. | yes",
            "Kyra.Kneeland55 | Kyra Kneeland | Email handle: kyrakneeland55@gmail.com. | yes",
            "Kyra.Kneeland55 | Kyra Kneeland | Private identifier string, no transcript meaning. | yes",
            "Zaden | Zaden | Name already supplied by the client. | yes",
            "Zaliah | Zaliah | Name already supplied by the client. | yes",
            "Tyshawn | Tyshawn | Name already supplied by the client. | yes",
            "Caseworker | Caseworker | Generic role, not a research entity. | yes",
            "Family to Family Program | Family to Family Program | A family-support initiative relevant to the program discussed. | yes",
            "Family to Family Program | Family to Family Program | Duplicate search finding for the same initiative. | yes",
        ])
        result = NS["_human_ai_agent_research_footer"](transcript, research)
        self.assertIn("Client spellings: Kneeland, Kyra, Zaden, Zaliah, Zaire, Tyshawn; I researched: Family to Family Program.", result)
        self.assertEqual(result.count("- Family to Family Program:"), 1)
        for unwanted in ("My spellings: None", "Kyra.Kneeland55", "gmail.com", "Caseworker:", "- Kneeland:", "Old model-written"):
            self.assertNotIn(unwanted, result)

    def test_worker_draft_charges_audio_minutes_plus_format_and_proofreading_costs_five(self):
        formatter_calls = {node.func.id for node in ast.walk(self.functions["_human_worker_format_ai_draft"]) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)}
        self.assertTrue({"_admin_guidelines_text", "_human_review_context", "_human_ai_agent_research", "_human_call_model_chain", "_human_worker_ai_draft_system", "_review_normalise_sentence_spacing", "_review_enforce_indent", "_human_ai_agent_research_footer"}.issubset(formatter_calls))
        route = self.functions["human_worker_ai_draft"]
        route_source = ast.unparse(route)
        self.assertIn("_human_worker_ai_draft_credit_cost", route_source)
        self.assertIn("plan_spend", route_source)
        self.assertIn("firestore.transactional", route_source)
        self.assertIn("worker_ai_draft", route_source)
        self.assertIn("format_version': 7", route_source)
        self.assertIn("credits_remaining", route_source)
        self.assertNotIn("charge_credits", route_source)

        proofread = self.functions["human_worker_ai_proofread_draft"]
        proofread_source = ast.unparse(proofread)
        self.assertIn("WORKER_DRAFT_PROOFREAD_MODEL_CHAIN", proofread_source)
        self.assertIn("_admin_guidelines_text", proofread_source)
        self.assertIn("_human_review_context", proofread_source)
        self.assertIn("_human_review_spelling_notes", proofread_source)
        self.assertIn("_human_review_full_audio_deepgram", proofread_source)
        self.assertIn("deepgram_comparison", proofread_source)
        self.assertIn("plan_spend", proofread_source)
        self.assertIn("credit_ledger", proofread_source)
        self.assertIn("source_sha256", proofread_source)
        self.assertIn("firestore.transactional", proofread_source)
        self.assertIn("credits_charged", proofread_source)
        self.assertIn("HUMAN_WORK_AI_PROOFREAD_CREDIT_COST", proofread_source)
        self.assertIn("plan_spend(profile_snapshot.to_dict() or {}, HUMAN_WORK_AI_PROOFREAD_CREDIT_COST)", proofread_source)
        self.assertNotIn("plan_spend(profile_snapshot.to_dict() or {}, 1)", proofread_source)
        self.assertNotIn("Claude Haiku", proofread_source)

        admin_proofread = ast.unparse(self.functions["human_admin_ai_review"])
        self.assertIn("HUMAN_WORK_AI_PROOFREAD_CREDIT_COST if is_human_subadmin", admin_proofread)
        self.assertIn("research_status = 'partial'", admin_proofread)
        self.assertIn("research_status = 'unavailable'", admin_proofread)
        self.assertNotIn("Required proper-noun web research did not complete", admin_proofread)
        self.assertIn("Claude Sonnet 5.5", admin_proofread)
        self.assertIn("fallback", admin_proofread)

    def test_worker_draft_cost_rounds_up_assigned_audio_minutes(self):
        function = self.functions["_human_worker_ai_draft_credit_cost"]
        namespace = {"math": math}
        exec(compile(ast.Module(body=[function], type_ignores=[]), "main.py", "exec"), namespace)
        cost = namespace["_human_worker_ai_draft_credit_cost"]
        self.assertEqual(cost({"seconds": 60}), {"audio_minutes": 1, "audio_credits": 1, "formatting_credits": 1, "total_credits": 2})
        self.assertEqual(cost({"seconds": 61}), {"audio_minutes": 2, "audio_credits": 2, "formatting_credits": 1, "total_credits": 3})
        self.assertEqual(cost({"seconds": 900}, {"start_seconds": 300, "end_seconds": 361}), {"audio_minutes": 2, "audio_credits": 2, "formatting_credits": 1, "total_credits": 3})

    def test_non_admin_job_serializer_hides_internal_model_metadata(self):
        serializer = ast.unparse(self.functions["_human_public_for"])
        self.assertIn("out.pop('ai_drafts', None)", serializer)
        self.assertIn("out['ai_draft'] = draft.get('text')", serializer)
        self.assertIn("out['ai_draft_proofread'] = proofread.get('text')", serializer)

    def test_worker_draft_prompt_requires_research_notes_and_handles_clear_speaker_corrections(self):
        function = self.functions["_human_worker_ai_draft_system"]
        prompt_text = "".join(node.value for node in ast.walk(function) if isinstance(node, ast.Constant) and isinstance(node.value, str))
        self.assertIn("Research Notes:", prompt_text)
        self.assertIn("I searched:", prompt_text)
        self.assertIn("Dublin Granville-East Dublin Granville Children's Close To Home.", prompt_text)
        self.assertIn("East Dublin Granville Children's Close To Home.", prompt_text)
        self.assertIn("Do not remove ordinary repetition", prompt_text)
        self.assertIn("non-semantic spoken fillers such as `um`, `uh`, or `you know`", prompt_text)
        self.assertIn("only if that name or term was actually spoken", prompt_text)
        self.assertIn("Use quotation marks only when quotation was dictated or to mark actual reported speech", prompt_text)

    def test_research_notes_retry_missing_grounding_and_fail_closed(self):
        class QuietLogger:
            def warning(self, *args, **kwargs):
                pass

        class QuietTime:
            @staticmethod
            def sleep(_seconds):
                pass

        class Response:
            status_code = 200
            text = ""

            def __init__(self, payload, status_code=200):
                self.payload = payload
                self.status_code = status_code

            def json(self):
                return self.payload

        class FakeRequests:
            responses = []
            calls = []

            @classmethod
            def post(cls, *args, **kwargs):
                cls.calls.append((args, kwargs))
                return cls.responses.pop(0)

        function = self.functions["_gemini_research_blocking"]
        namespace = {"GEMINI_API_KEY": "test-key", "requests": FakeRequests, "logger": QuietLogger(), "time": QuietTime()}
        exec(compile(ast.Module(body=[function], type_ignores=[]), "main.py", "exec"), namespace)
        ungrounded = Response({"candidates": [{"content": {"parts": [{"text": "A likely result."}]}}]})
        grounded_response = Response({"candidates": [{
            "content": {"parts": [{"text": "Summit Psych | Summit Psych | provider | yes"}]},
            "groundingMetadata": {
                "webSearchQueries": ["Summit Psych Ohio provider"],
                "groundingChunks": [{"web": {"title": "Summit Psych", "uri": "https://example.test"}}],
            },
        }]})
        FakeRequests.responses = [ungrounded, ungrounded, grounded_response]
        FakeRequests.calls = []
        grounded = namespace["_gemini_research_blocking"]("prompt")
        self.assertIn("ACTUAL GOOGLE SEARCH QUERIES:\n- Summit Psych Ohio provider", grounded)
        self.assertIn("ACTUAL SEARCH SOURCES:\n- Summit Psych: https://example.test", grounded)
        self.assertEqual(len(FakeRequests.calls), 3)
        self.assertEqual(FakeRequests.calls[-1][1]["timeout"], (10, 150))

        FakeRequests.responses = [Response({"candidates": [{"content": {"parts": [{"text": "NO_SEARCHED_TERMS"}]}}]})]
        self.assertEqual(namespace["_gemini_research_blocking"]("prompt"), "NO_SEARCHED_TERMS")

        FakeRequests.responses = [ungrounded, ungrounded, ungrounded]
        with self.assertRaisesRegex(RuntimeError, "no grounded queries or sources"):
            namespace["_gemini_research_blocking"]("prompt")
        self.assertEqual(len(FakeRequests.responses), 0)


class TemplateAgentGuidelinesAndPrivateReferences(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source_path = Path(__file__).resolve().parents[1].joinpath("main.py")
        cls.source = source_path.read_text(encoding="utf-8")
        cls.tree = ast.parse(cls.source)
        cls.functions = {
            node.name: node for node in cls.tree.body
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
        }
        agent_node = next(
            node for node in cls.tree.body if isinstance(node, ast.Assign)
            and any(getattr(target, "id", None) == "HUMAN_AI_AGENTS" for target in node.targets)
        )
        correction_node = next(
            node for node in cls.tree.body if isinstance(node, ast.Assign)
            and any(getattr(target, "id", None) == "HUMAN_GENERAL_SELF_CORRECTION_GUIDANCE" for target in node.targets)
        )
        namespace = {
            "os": __import__("os"), "HUMAN_AI_AGENTS": ast.literal_eval(agent_node.value),
            "HUMAN_GENERAL_SELF_CORRECTION_GUIDANCE": ast.literal_eval(correction_node.value),
            "__file__": str(source_path),
        }
        selected = [
            cls.functions["_human_template_agent_guidelines"],
            cls.functions["_human_ai_agent_system"],
        ]
        exec(compile(ast.Module(body=selected, type_ignores=[]), str(source_path), "exec"), namespace)
        cls.helpers = namespace

    def test_full_template_rule_file_is_loaded_and_uses_researched_label(self):
        rules = self.helpers["_human_template_agent_guidelines"]()
        for phrase in (
            "Children Services", "walk-through", "Client spellings:", "My spellings:",
            "I researched:", "Research Notes:", "Jeremy, J-E-R-A-M-I-E",
            "TEMPLATE_JOB_BLOCKED_LETTER", "Never insert `[dictation ends here]`",
            "Quantifying numbers", "5'10\"", "straight ASCII",
        ):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, rules)
        self.assertNotIn("I searched:", rules)

    def test_template_prompt_contains_permanent_and_current_job_rules(self):
        prompt = self.helpers["_human_ai_agent_system"](
            "template-claude", "transcription and formatting", "older general instructions", "job files",
            "Keep client-confirmed name spellings and preserve dictated paragraphing.",
        )
        self.assertIn("PERMANENT TEMPLATE-JOB GUIDELINES", prompt)
        self.assertIn("I researched:", prompt)
        self.assertIn("never `I searched:`", prompt)
        self.assertIn("Below are text-specific guidelines", prompt)
        self.assertIn("Keep client-confirmed name spellings", prompt)
        self.assertIn("job files", prompt)
        self.assertIn("TEMPLATE_JOB_BLOCKED_LETTER", prompt)

    def test_other_agent_prompt_does_not_receive_template_rules(self):
        prompt = self.helpers["_human_ai_agent_system"](
            "general-gpt", "transcription and formatting", "general instructions", "job files",
        )
        self.assertNotIn("PERMANENT TEMPLATE-JOB GUIDELINES", prompt)
        self.assertNotIn("CURRENT JOB TEMPLATE-SPECIFIC GUIDELINES", prompt)

    def test_general_agent_and_reviewer_only_receive_clear_self_correction_rule(self):
        general_prompt = self.helpers["_human_ai_agent_system"](
            "general-gpt", "transcription and formatting", "general instructions", "job files",
        )
        template_prompt = self.helpers["_human_ai_agent_system"](
            "template-claude", "transcription and formatting", "general instructions", "job files",
        )
        self.assertIn("Dublin Granville— East Dublin Granville", general_prompt)
        self.assertIn("Preserve all other wording, order, grammar, and repetition", general_prompt)
        self.assertNotIn("GENERAL-JOBS SPEAKER SELF-CORRECTIONS", template_prompt)
        reviewer = ast.unparse(self.functions["human_admin_ai_review"])
        self.assertIn("HUMAN_GENERAL_SELF_CORRECTION_GUIDANCE if job_type == 'general_job'", reviewer)

    def test_current_job_notes_and_uploaded_references_are_saved_privately_per_run(self):
        assign = ast.unparse(next(node for node in self.tree.body if isinstance(node, ast.AsyncFunctionDef) and node.name == "human_admin_assign_ai_agent"))
        runner = ast.unparse(self.functions["_human_run_ai_agent"])
        generate = ast.unparse(self.functions["_human_ai_agent_generate"])
        context = ast.unparse(self.functions["_human_review_context"])
        self.assertIn("request.form()", assign)
        self.assertIn("job_specific_guidelines", assign)
        self.assertIn("reference_files", assign)
        self.assertIn("template_reference_files", assign)
        self.assertIn("ai-agent-runs/", assign)
        self.assertIn("run_snapshot", runner)
        self.assertIn("template_job_guidelines", runner)
        self.assertIn("private_reference_files", runner)
        self.assertIn("private_attachments=private_reference_files", generate)
        self.assertIn("human-workflow/{job_id}/ai-agent-runs/", context)
        self.assertIn("TEMPLATE_JOB_BLOCKED_LETTER", runner)

    def test_template_agent_never_adds_dictation_ending_marker_to_parts(self):
        prompt = self.helpers["_human_ai_agent_system"](
            "template-claude", "transcription and formatting", "", "", "",
        )
        self.assertIn("never add the marker `[dictation ends here]`", prompt)
        worker_prompt = ast.unparse(self.functions["_human_worker_ai_draft_system"])
        self.assertIn("Never add `[dictation ends here]`", worker_prompt)

