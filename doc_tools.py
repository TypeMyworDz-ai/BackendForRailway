"""Free document conversion used by the public Tools pages and by PDF Jobs.

Nothing here touches a database or keeps a file. Every conversion works on
bytes in a temporary folder that is removed before the function returns.

Word to PDF needs LibreOffice (the `soffice` program). The Docker image tries
to install it; if it is missing, the functions raise ConversionError with a
plain message instead of crashing the app.
"""
import io
import os
import re
import shutil
import subprocess
import tempfile
import threading
import zipfile

OFFICE_EXTENSIONS = {"docx", "doc", "rtf", "odt", "txt"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff", "gif"}
MAX_PDF_PAGES_FOR_IMAGES = 40
MAX_PDF_PAGES_FOR_WORD = 60

# LibreOffice is heavy. Two at once is plenty for a free tool.
_OFFICE_SLOTS = threading.BoundedSemaphore(2)


class ConversionError(Exception):
    """A problem worth showing to the person who uploaded the file."""


def file_extension(name):
    name = os.path.basename(name or "")
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def safe_stem(name, default="document"):
    stem = os.path.splitext(os.path.basename(name or ""))[0]
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "", stem).strip(" .")[:80]
    return stem or default


def soffice_path():
    return shutil.which("soffice") or shutil.which("libreoffice")


def office_to_pdf_bytes(raw, filename="document.docx", timeout=120):
    """Word (or similar) to PDF using LibreOffice in a throwaway profile."""
    binary = soffice_path()
    if not binary:
        raise ConversionError("Word conversion is not available on this server right now. Please try again later.")
    ext = file_extension(filename)
    if ext not in OFFICE_EXTENSIONS:
        raise ConversionError("That file type is not a Word document.")
    work = tempfile.mkdtemp(prefix="tm-office-")
    try:
        source = os.path.join(work, f"input.{ext}")
        with open(source, "wb") as handle:
            handle.write(raw)
        profile = os.path.join(work, "profile")
        command = [
            binary, "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nofirststartwizard",
            f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf:writer_pdf_Export", "--outdir", work, source,
        ]
        env = dict(os.environ, HOME=work)
        with _OFFICE_SLOTS:
            try:
                subprocess.run(command, cwd=work, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=timeout, check=False)
            except subprocess.TimeoutExpired as exc:
                raise ConversionError("The document took too long to convert. Try a smaller file.") from exc
        produced = os.path.join(work, "input.pdf")
        if not os.path.exists(produced):
            raise ConversionError("That document could not be converted. It may be damaged or password protected.")
        with open(produced, "rb") as handle:
            return handle.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _render_pdf_pages(raw, image_format, max_pages, scale=2.0):
    import pypdfium2 as pdfium

    try:
        document = pdfium.PdfDocument(raw)
    except Exception as exc:
        raise ConversionError("That PDF could not be opened. It may be damaged or password protected.") from exc
    try:
        count = len(document)
        if count < 1:
            raise ConversionError("That PDF has no pages.")
        if count > max_pages:
            raise ConversionError(f"That PDF has {count} pages. The free tool handles up to {max_pages} pages at a time.")
        pages = []
        for index in range(count):
            page = document[index]
            width, height = page.get_size()
            fit = min(scale, (16_000_000 / max(float(width * height), 1.0)) ** 0.5)
            image = page.render(scale=max(0.5, fit)).to_pil().convert("RGB")
            out = io.BytesIO()
            if image_format == "png":
                image.save(out, format="PNG", optimize=True)
            else:
                image.save(out, format="JPEG", quality=92, optimize=True)
            pages.append(out.getvalue())
        return pages
    finally:
        try:
            document.close()
        except Exception:
            pass


def pdf_to_image_list(raw, image_format="jpg", max_pages=MAX_PDF_PAGES_FOR_IMAGES):
    return _render_pdf_pages(raw, "png" if image_format == "png" else "jpg", max_pages)


def pdf_to_docx_bytes(raw):
    """PDF to an editable Word file. Scanned PDFs have no text to carry over."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(raw))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                raise ConversionError("That PDF is password protected.")
        pages = len(reader.pages)
        if pages > MAX_PDF_PAGES_FOR_WORD:
            raise ConversionError(f"That PDF has {pages} pages. The free tool handles up to {MAX_PDF_PAGES_FOR_WORD} pages at a time.")
        has_text = any((page.extract_text() or "").strip() for page in reader.pages[:10])
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("That PDF could not be read. It may be damaged.") from exc
    if not has_text:
        raise ConversionError("This PDF looks like a scan, so there is no text to put into Word. Convert it to images instead, or transcribe the pages with TypeMyworDz.")
    work = tempfile.mkdtemp(prefix="tm-pdf-")
    try:
        source = os.path.join(work, "input.pdf")
        target = os.path.join(work, "output.docx")
        with open(source, "wb") as handle:
            handle.write(raw)
        try:
            from pdf2docx import Converter

            converter = Converter(source)
            try:
                converter.convert(target)
            finally:
                converter.close()
        except ImportError:
            return _pdf_text_to_docx(reader)
        except Exception:
            return _pdf_text_to_docx(reader)
        if not os.path.exists(target):
            return _pdf_text_to_docx(reader)
        with open(target, "rb") as handle:
            return handle.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _pdf_text_to_docx(reader):
    """Plain fallback: the text of every page, one page after another."""
    from docx import Document

    document = Document()
    for number, page in enumerate(reader.pages):
        if number:
            document.add_page_break()
        text = page.extract_text() or ""
        for block in re.split(r"\n\s*\n", text):
            block = " ".join(block.split())
            if block:
                document.add_paragraph(block)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def images_to_pdf_bytes(image_blobs):
    from PIL import Image, ImageOps

    frames = []
    for blob in image_blobs:
        try:
            image = Image.open(io.BytesIO(blob))
            image.load()
        except Exception as exc:
            raise ConversionError("One of the pictures could not be read.") from exc
        frames.append(ImageOps.exif_transpose(image).convert("RGB"))
    if not frames:
        raise ConversionError("Choose at least one picture.")
    out = io.BytesIO()
    frames[0].save(out, format="PDF", save_all=True, append_images=frames[1:], resolution=150.0)
    return out.getvalue()


def image_convert_bytes(raw, image_format):
    from PIL import Image, ImageOps

    try:
        image = Image.open(io.BytesIO(raw))
        image.load()
    except Exception as exc:
        raise ConversionError("That picture could not be read.") from exc
    image = ImageOps.exif_transpose(image)
    out = io.BytesIO()
    if image_format == "png":
        image.convert("RGBA" if image.mode in ("RGBA", "LA", "P") else "RGB").save(out, format="PNG", optimize=True)
    else:
        image.convert("RGB").save(out, format="JPEG", quality=92, optimize=True)
    return out.getvalue()


def _zip_named(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files:
            archive.writestr(name, data)
    return out.getvalue()


TARGETS = {"pdf", "docx", "png", "jpg"}


def convert(filename, raw, target):
    """Return (download_name, content_type, bytes) or raise ConversionError."""
    target = (target or "").lower().replace("jpeg", "jpg")
    if target not in TARGETS:
        raise ConversionError("Choose what to convert the file into.")
    ext = file_extension(filename)
    stem = safe_stem(filename)
    is_pdf = ext == "pdf" or raw[:5] == b"%PDF-"
    if is_pdf:
        if target == "docx":
            return f"{stem}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", pdf_to_docx_bytes(raw)
        if target in ("png", "jpg"):
            pages = pdf_to_image_list(raw, target)
            mime = "image/png" if target == "png" else "image/jpeg"
            if len(pages) == 1:
                return f"{stem}.{target}", mime, pages[0]
            return f"{stem}-{target}-pages.zip", "application/zip", _zip_named([(f"{stem}-page-{i + 1:02d}.{target}", data) for i, data in enumerate(pages)])
        raise ConversionError("That file is already a PDF.")
    if ext in OFFICE_EXTENSIONS:
        pdf = office_to_pdf_bytes(raw, filename)
        if target == "pdf":
            return f"{stem}.pdf", "application/pdf", pdf
        if target in ("png", "jpg"):
            pages = pdf_to_image_list(pdf, target)
            mime = "image/png" if target == "png" else "image/jpeg"
            if len(pages) == 1:
                return f"{stem}.{target}", mime, pages[0]
            return f"{stem}-{target}-pages.zip", "application/zip", _zip_named([(f"{stem}-page-{i + 1:02d}.{target}", data) for i, data in enumerate(pages)])
        raise ConversionError("That file is already a Word document.")
    if ext in IMAGE_EXTENSIONS:
        if target == "pdf":
            return f"{stem}.pdf", "application/pdf", images_to_pdf_bytes([raw])
        if target in ("png", "jpg"):
            return f"{stem}.{target}", "image/png" if target == "png" else "image/jpeg", image_convert_bytes(raw, target)
        raise ConversionError("A picture can be turned into a PDF, PNG or JPG.")
    raise ConversionError("That file type is not supported. Use a Word document, a PDF or a picture.")


# ---------------------------------------------------------------------------
# Merge, split and compress PDFs, and the audio converter
# ---------------------------------------------------------------------------
MAX_MERGE_FILES = 12
MAX_SPLIT_PAGES = 200
AUDIO_TARGETS = {
    "mp3": ("audio/mpeg", ["-c:a", "libmp3lame"]),
    "wav": ("audio/wav", ["-c:a", "pcm_s16le"]),
    "m4a": ("audio/mp4", ["-c:a", "aac"]),
    "ogg": ("audio/ogg", ["-c:a", "libvorbis"]),
    "flac": ("audio/flac", ["-c:a", "flac"]),
}
AUDIO_BITRATES = {"32", "48", "64", "96", "128", "192", "256"}


def _open_reader(raw, what="PDF"):
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ConversionError(f"That {what} is password protected. Remove the password first.")
            except ConversionError:
                raise
            except Exception as exc:
                raise ConversionError(f"That {what} is password protected. Remove the password first.") from exc
        len(reader.pages)
        return reader
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"That {what} could not be opened. It may be damaged.") from exc


def merge_pdfs(files):
    """files: list of (filename, raw). PDFs and pictures are accepted, in order."""
    from pypdf import PdfWriter

    if len(files) < 2:
        raise ConversionError("Add at least two files to merge.")
    if len(files) > MAX_MERGE_FILES:
        raise ConversionError(f"You can merge up to {MAX_MERGE_FILES} files at a time.")
    writer = PdfWriter()
    for name, raw in files:
        ext = file_extension(name)
        if ext in IMAGE_EXTENSIONS:
            raw = images_to_pdf_bytes([raw])
        elif ext in OFFICE_EXTENSIONS:
            raw = office_to_pdf_bytes(raw, name)
        elif not (ext == "pdf" or raw[:5] == b"%PDF-"):
            raise ConversionError(f"{os.path.basename(name)} is not a PDF or a picture.")
        for page in _open_reader(raw).pages:
            writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    return "merged.pdf", "application/pdf", out.getvalue()


def parse_page_ranges(text, total):
    """'1-3, 5, 8-' -> [(1,3),(5,5),(8,total)]. Raises ConversionError."""
    ranges = []
    for part in re.split(r"[,;\s]+", (text or "").strip()):
        if not part:
            continue
        match = re.fullmatch(r"(\d*)\s*-\s*(\d*)|(\d+)", part)
        if not match:
            raise ConversionError(f"'{part}' is not a page range. Use something like 1-3, 5, 8-10.")
        if match.group(3):
            start = end = int(match.group(3))
        else:
            start = int(match.group(1) or 1)
            end = int(match.group(2) or total)
        if start < 1 or end < start or end > total:
            raise ConversionError(f"Pages {part} are outside this PDF, which has {total} pages.")
        ranges.append((start, end))
    if not ranges:
        raise ConversionError("Type the pages you want, for example 1-3, 5.")
    return ranges


def _pdf_from_pages(reader, pages):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for number in pages:
        writer.add_page(reader.pages[number - 1])
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def split_pdf(filename, raw, mode="ranges", ranges_text=""):
    reader = _open_reader(raw)
    total = len(reader.pages)
    stem = safe_stem(filename)
    if total > MAX_SPLIT_PAGES:
        raise ConversionError(f"That PDF has {total} pages. The free tool handles up to {MAX_SPLIT_PAGES} pages.")
    if mode == "each":
        files = [(f"{stem}-page-{n:03d}.pdf", _pdf_from_pages(reader, [n])) for n in range(1, total + 1)]
        return f"{stem}-pages.zip", "application/zip", _zip_named(files)
    ranges = parse_page_ranges(ranges_text, total)
    if mode == "extract":
        pages = [n for start, end in ranges for n in range(start, end + 1)]
        return f"{stem}-extract.pdf", "application/pdf", _pdf_from_pages(reader, pages)
    if len(ranges) == 1:
        start, end = ranges[0]
        return f"{stem}-{start}-{end}.pdf", "application/pdf", _pdf_from_pages(reader, list(range(start, end + 1)))
    files = [(f"{stem}-{start}-{end}.pdf", _pdf_from_pages(reader, list(range(start, end + 1)))) for start, end in ranges]
    return f"{stem}-split.zip", "application/zip", _zip_named(files)


_COMPRESS_LEVELS = {"low": (82, 2400), "recommended": (62, 1700), "strong": (42, 1200)}


def compress_pdf(filename, raw, level="recommended"):
    """Returns (name, content_type, bytes, original_size). Never returns a bigger file."""
    from pypdf import PdfWriter
    from PIL import Image

    quality, longest = _COMPRESS_LEVELS.get(level, _COMPRESS_LEVELS["recommended"])
    reader = _open_reader(raw)
    writer = PdfWriter(clone_from=reader)
    for page in writer.pages:
        try:
            images = list(page.images)
        except Exception:
            images = []
        for image in images:
            try:
                pil = image.image
                if pil is None:
                    continue
                if max(pil.size) > longest:
                    scale = longest / float(max(pil.size))
                    pil = pil.resize((max(1, int(pil.width * scale)), max(1, int(pil.height * scale))), Image.LANCZOS)
                if pil.mode not in ("RGB", "L"):
                    pil = pil.convert("RGB")
                image.replace(pil, quality=quality)
            except Exception:
                continue
        try:
            page.compress_content_streams()
        except Exception:
            pass
    try:
        writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    except Exception:
        pass
    out = io.BytesIO()
    writer.write(out)
    data = out.getvalue()
    if len(data) >= len(raw):
        data = raw
    return f"{safe_stem(filename)}-compressed.pdf", "application/pdf", data, len(raw)


def audio_convert(filename, raw, target="mp3", bitrate="96", mono=False, timeout=180, trim_silence=False):
    """Any audio or video file in, a smaller or different audio file out, via ffmpeg."""
    target = (target or "mp3").lower()
    if target not in AUDIO_TARGETS:
        raise ConversionError("Choose MP3, WAV, M4A, OGG or FLAC.")
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise ConversionError("The audio converter is temporarily unavailable. Please try again later.")
    bitrate = str(bitrate) if str(bitrate) in AUDIO_BITRATES else "96"
    mime, codec = AUDIO_TARGETS[target]
    ext_in = re.sub(r"[^a-z0-9]", "", file_extension(filename)) or "bin"
    with tempfile.TemporaryDirectory() as folder:
        src = os.path.join(folder, f"in.{ext_in}")
        dst = os.path.join(folder, f"out.{target}")
        with open(src, "wb") as handle:
            handle.write(raw)
        command = [ffmpeg, "-y", "-hide_banner", "-loglevel", "error", "-i", src, "-vn", "-map_metadata", "-1"] + codec
        if target in ("mp3", "m4a", "ogg"):
            command += ["-b:a", f"{bitrate}k"]
        if mono:
            command += ["-ac", "1"]
        if trim_silence:
            # Cut the lead-in and shorten every pause longer than 1.2 seconds to 0.4 seconds.
            command += ["-af", "silenceremove=start_periods=1:start_threshold=-45dB:start_silence=0.2:stop_periods=-1:stop_duration=1.2:stop_threshold=-45dB:stop_silence=0.4"]
        command.append(dst)
        try:
            result = subprocess.run(command, capture_output=True, timeout=timeout)
        except subprocess.TimeoutExpired as exc:
            raise ConversionError("That file took too long to convert. Try a shorter recording.") from exc
        if result.returncode != 0 or not os.path.exists(dst) or os.path.getsize(dst) == 0:
            raise ConversionError("That file could not be read as audio. Try an MP3, WAV, M4A, WEBM or MP4 file.")
        with open(dst, "rb") as handle:
            data = handle.read()
    return f"{safe_stem(filename, 'audio')}.{target}", mime, data
